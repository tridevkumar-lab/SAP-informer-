"""Small local SQLite full-text knowledge store for the starter release."""
from pathlib import Path
import re
import sqlite3
from typing import Iterable


def connect(db_path: str) -> sqlite3.Connection:
    Path(db_path).parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.execute("""CREATE TABLE IF NOT EXISTS chunks (
        id INTEGER PRIMARY KEY, source TEXT NOT NULL, chunk_no INTEGER NOT NULL,
        content TEXT NOT NULL, UNIQUE(source, chunk_no))""")
    conn.execute("""CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(
        content, source UNINDEXED)""")
    conn.commit()
    return conn


def split_text(text: str, size: int = 1000, overlap: int = 150) -> list[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + size // 2, end)
            if boundary > start:
                end = boundary
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def upsert_document(db_path: str, source: str, text: str) -> int:
    chunks = split_text(text)
    with connect(db_path) as conn:
        old = conn.execute("SELECT id FROM chunks WHERE source = ?", (source,)).fetchall()
        for row in old:
            conn.execute("DELETE FROM chunks_fts WHERE rowid = ?", (row["id"],))
        conn.execute("DELETE FROM chunks WHERE source = ?", (source,))
        for number, content in enumerate(chunks, start=1):
            cur = conn.execute(
                "INSERT INTO chunks(source, chunk_no, content) VALUES (?, ?, ?)",
                (source, number, content),
            )
            conn.execute(
                "INSERT INTO chunks_fts(rowid, content, source) VALUES (?, ?, ?)",
                (cur.lastrowid, content, source),
            )
    return len(chunks)


def search(db_path: str, query: str, limit: int = 5) -> list[dict[str, str]]:
    # Quote terms before passing them to the FTS parser; user input is not SQL.
    terms = re.findall(r"[\w.-]+", query, flags=re.UNICODE)
    if not terms:
        return []
    match = " OR ".join('"' + term.replace('"', '""') + '"' for term in terms[:20])
    with connect(db_path) as conn:
        rows = conn.execute(
            """SELECT c.source, c.content FROM chunks_fts f
               JOIN chunks c ON c.id = f.rowid
               WHERE chunks_fts MATCH ? ORDER BY bm25(chunks_fts) LIMIT ?""",
            (match, limit),
        ).fetchall()
    return [{"source": row["source"], "content": row["content"]} for row in rows]


def all_sources(db_path: str) -> list[str]:
    with connect(db_path) as conn:
        rows = conn.execute("SELECT DISTINCT source FROM chunks ORDER BY source").fetchall()
    return [row["source"] for row in rows]


def _extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".log"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines)
    raise ValueError(f"Unsupported file type: {suffix}")


def ingest_folder(db_path: str, folder: str) -> dict[str, int]:
    root = Path(folder)
    root.mkdir(parents=True, exist_ok=True)
    supported = {".txt", ".md", ".csv", ".log", ".pdf", ".docx", ".xlsx"}
    result = {}
    for path in root.rglob("*"):
        if path.is_file() and path.suffix.lower() in supported:
            result[path.name] = upsert_document(db_path, path.name, _extract(path))
    return result
